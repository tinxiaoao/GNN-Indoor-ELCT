"""Compare saved outputs, keeping migration and scientific criteria separate.

Reference tables are read only by this explicit evaluator, after extraction.
The retained legacy and strict evaluators are called without changing their
default IoU matching threshold or exact-value comparison rules.
"""
from __future__ import annotations

import argparse
from collections import Counter
import csv
import json
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from geo_elct.evaluation.compare_elct_tables import run_comparison
from geo_elct.evaluation.evaluate_reference_faithful_tables import evaluate_tables


def dump(path: Path, data: Any) -> None:
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def write_csv(path: Path, rows: list[dict], fields=None) -> None:
    fields = fields or (list(rows[0]) if rows else [])
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def workbook(path: Path) -> dict:
    book = load_workbook(path, read_only=True, data_only=False)
    data = {sheet.title: [list(row) for row in sheet.iter_rows(values_only=True)] for sheet in book.worksheets}
    book.close()
    return data


def full_cells(reference: Path, candidate: Path) -> dict:
    ref, cand = workbook(reference), workbook(candidate)
    differences, schemas = [], []
    for sheet in sorted(set(ref) | set(cand)):
        a, b = ref.get(sheet, []), cand.get(sheet, [])
        schemas.append({"sheet": sheet, "reference_rows": max(len(a) - 1, 0),
                        "candidate_rows": max(len(b) - 1, 0),
                        "reference_headers": a[0] if a else [],
                        "candidate_headers": b[0] if b else [],
                        "headers_identical": bool(a and b and a[0] == b[0])})
        for row in range(max(len(a), len(b))):
            ra, rb = a[row] if row < len(a) else [], b[row] if row < len(b) else []
            for col in range(max(len(ra), len(rb))):
                va, vb = ra[col] if col < len(ra) else None, rb[col] if col < len(rb) else None
                if va != vb:
                    differences.append({"sheet": sheet, "row": row + 1, "column": col + 1,
                                        "header": a[0][col] if a and col < len(a[0]) else None,
                                        "reference_value": va, "candidate_value": vb})
    return {"identical": ref == cand, "schemas": schemas, "cell_differences": differences}


def boundary_records(folder: Path, mapping=None) -> Counter:
    stats = json.loads((folder / "room_stats.json").read_text(encoding="utf-8"))
    records = Counter()
    for room_id, data in stats.items():
        rid = int(room_id)
        rid = mapping.get(rid, f"unmatched_candidate:{rid}") if mapping is not None else rid
        for field in ("walls", "openings"):
            for segment in data.get(field, []):
                records[(rid, field, *segment)] += 1
    return records


def counter_rows(records: Counter) -> list[dict]:
    return [{"room_id": row[0], "kind": row[1], "record": list(row[2:]), "multiplicity": count}
            for row, count in sorted(records.items(), key=lambda item: str(item[0]))]


def spatial_comparison(reference: Path, candidate: Path, mapping: dict, scores: dict,
                       tx_room_sets: list[list[int]]) -> dict:
    ref, cand = np.load(reference / "region_id_map.npy"), np.load(candidate / "region_id_map.npy")
    ref_ids, cand_ids = set(int(x) for x in np.unique(ref) if x > 0), set(int(x) for x in np.unique(cand) if x > 0)
    rows = []
    for cid, rid in sorted(mapping.items()):
        ry, rx = np.where(ref == rid)
        cy, cx = np.where(cand == cid)
        def geometry(xs, ys):
            return {"area_px": int(len(xs)), "centroid_xy": [float(xs.mean()), float(ys.mean())],
                    "bbox_xyxy": [int(xs.min()), int(ys.min()), int(xs.max()), int(ys.max())]}
        rows.append({"reference_room_id": rid, "candidate_room_id": cid,
                     "iou": scores.get(str(cid), scores.get(cid)),
                     "identical_mask": bool(np.array_equal(ref == rid, cand == cid)),
                     "reference_geometry": geometry(rx, ry), "candidate_geometry": geometry(cx, cy)})
    reverse = {rid: cid for cid, rid in mapping.items()}
    tx = [{"case_index": index + 1, "reference_tx_room_ids": ids,
           "candidate_spatial_tx_room_ids": [reverse.get(rid) for rid in ids],
           "ids_unchanged": all(reverse.get(rid) == rid for rid in ids),
           "all_tx_masks_identical": all(rid in reverse and np.array_equal(ref == rid, cand == reverse[rid]) for rid in ids)}
          for index, ids in enumerate(tx_room_sets)]
    return {"region_map_identical": bool(np.array_equal(ref, cand)), "rooms": rows, "tx": tx,
            "unmatched_reference_rooms": sorted(ref_ids - set(mapping.values())),
            "unmatched_candidate_rooms": sorted(cand_ids - set(mapping)),
            "all_matched_masks_identical": all(row["identical_mask"] for row in rows),
            "all_room_ids_unchanged": all(cid == rid for cid, rid in mapping.items())}


def compare_outputs(reference: Path, candidate: Path, out: Path,
                    tx_room_sets: list[list[int]] | None = None) -> dict:
    reference, candidate, out = reference.resolve(), candidate.resolve(), out.resolve()
    out.mkdir(parents=True, exist_ok=False)
    strict = evaluate_tables(reference, candidate, out / "strict_tables")
    legacy = run_comparison(reference, candidate, out / "legacy_relation")
    mapping = {int(cid): int(rid) for cid, rid in strict["candidate_to_clean_room_mapping"].items()}
    cells = {name: full_cells(reference / filename, candidate / filename)
             for name, filename in [("room", "debug_roomself.xlsx"), ("edge", "debug_topology_edges.xlsx")]}
    dump(out / "all_workbook_cells.json", cells)
    spatial = spatial_comparison(reference, candidate, mapping, strict["room_match_scores"], tx_room_sets or [])
    dump(out / "spatial_and_tx.json", spatial)
    raw = {}
    for filename in ("rooms.json", "edges.json", "room_stats.json"):
        a = json.loads((reference / filename).read_text(encoding="utf-8"))
        b = json.loads((candidate / filename).read_text(encoding="utf-8"))
        raw[filename] = {"identical": a == b}
        if a != b:
            raw[filename].update(reference=a, candidate=b)
    dump(out / "raw_records.json", raw)
    ref_bounds, cand_bounds = boundary_records(reference), boundary_records(candidate, mapping)
    boundary = {"identical_spatially_mapped_records": ref_bounds == cand_bounds,
                "reference_records": counter_rows(ref_bounds), "candidate_records": counter_rows(cand_bounds),
                "removed": counter_rows(ref_bounds - cand_bounds), "added": counter_rows(cand_bounds - ref_bounds)}
    dump(out / "external_walls_and_openings.json", boundary)
    runtime_file = candidate / "runtime_validation.json"
    runtime = json.loads(runtime_file.read_text(encoding="utf-8")) if runtime_file.exists() else {}
    topology_file = candidate / "topology_summary.json"
    topology = json.loads(topology_file.read_text(encoding="utf-8")) if topology_file.exists() else {}
    summary = {"raw_room_workbook_identical": cells["room"]["identical"],
               "raw_edge_workbook_identical": cells["edge"]["identical"],
               "raw_records_identical": all(item["identical"] for item in raw.values()),
               "room_map_identical": spatial["region_map_identical"],
               "mask_matching_method": strict["room_matching_method"],
               "all_room_ids_unchanged": spatial["all_room_ids_unchanged"],
               "unmatched_reference_rooms": spatial["unmatched_reference_rooms"],
               "unmatched_candidate_rooms": spatial["unmatched_candidate_rooms"],
               "all_tx_masks_identical": all(row["all_tx_masks_identical"] for row in spatial["tx"]),
               "external_boundary_records_identical": boundary["identical_spatially_mapped_records"],
               "removed_external_opening_records": sum(row["multiplicity"] for row in boundary["removed"] if row["kind"] == "openings"),
               "legacy_topology_relation_preserved": legacy["topology_relation_preserved"],
               "legacy_door_F1": legacy.get("door_edge_F1"),
               "strict_internal_topology_exact": strict["topology_exact_match"],
               "strict_room_table_exact": strict["room_table_exact_match"],
               "strict_edge_table_exact": strict["edge_table_exact_match"],
               "strict_room_differences": strict["room_table_difference_count"],
               "strict_edge_differences": strict["edge_table_difference_count"],
               "runtime_reference_free_passed": runtime.get("passed"),
               "runtime_topology_summary_valid": topology.get("valid_topology")}
    dump(out / "separate_criteria_summary.json", summary)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--reference", type=Path, required=True)
    parser.add_argument("--candidate", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--tx-room-sets", type=Path, help="JSON list of reference-ID lists; spatial mapping is retained")
    args = parser.parse_args()
    tx = json.loads(args.tx_room_sets.read_text(encoding="utf-8")) if args.tx_room_sets else []
    result = compare_outputs(args.reference, args.candidate, args.out, tx)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
